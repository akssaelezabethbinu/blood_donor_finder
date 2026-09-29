from flask import Flask, render_template, request, redirect, session
import sqlite3
import os
from openpyxl import load_workbook

app = Flask(__name__)

app.secret_key = "blooddonor123"


# =========================
# CREATE DATABASE
# =========================

def create_db():

    conn = sqlite3.connect("donors.db")
    cursor = conn.cursor()

    # Donor Table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS donors(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT,
            age INTEGER,
            blood_group TEXT,
            phone TEXT,
            location TEXT
        )
    """)

    # Emergency Request Table
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS requests(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_name TEXT,
            blood_group TEXT,
            hospital TEXT,
            phone TEXT,
            priority TEXT
        )
    """)

    conn.commit()
    conn.close()


create_db()


# =========================
# HOME PAGE
# =========================

@app.route('/')
def home():
    return render_template("index.html")


# =========================
# REGISTER DONOR
# =========================

@app.route('/register', methods=['GET', 'POST'])
def register():

    if request.method == 'POST':

        name = request.form['name']
        age = request.form['age']
        blood_group = request.form['blood_group']
        phone = request.form['phone']
        location = request.form['location']

        conn = sqlite3.connect("donors.db")
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO donors
            (name, age, blood_group, phone, location)
            VALUES (?, ?, ?, ?, ?)
        """, (name, age, blood_group, phone, location))

        conn.commit()
        conn.close()

        return "Donor Registered Successfully!"

    return render_template("register.html")


# =========================
# FIND DONOR
# =========================

@app.route('/search', methods=['GET', 'POST'])
def search_page():

    donors = []

    if request.method == "POST":

        blood_group = request.form['blood_group']
        location = request.form['location']

        conn = sqlite3.connect("donors.db")
        cursor = conn.cursor()

        cursor.execute("""
            SELECT * FROM donors
            WHERE blood_group = ? AND location = ?
        """, (blood_group, location))

        donors = cursor.fetchall()

        conn.close()

    return render_template("search.html", donors=donors)


# =========================
# EMERGENCY REQUEST
# =========================

def calculate_priority(blood_group):

    rare_blood = ["O-", "AB-", "B-"]

    if blood_group in rare_blood:
        return "Critical 🔴"

    elif blood_group in ["A-", "O+"]:
        return "High 🟠"

    else:
        return "Normal 🟢"


@app.route('/request', methods=['GET', 'POST'])
def request_page():

    if request.method == "POST":

        patient_name = request.form['patient_name']
        blood_group = request.form['blood_group']
        hospital = request.form['hospital']
        phone = request.form['phone']

        priority = calculate_priority(blood_group)

        conn = sqlite3.connect("donors.db")
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO requests
            (patient_name, blood_group, hospital, phone, priority)
            VALUES (?, ?, ?, ?, ?)
        """, (
            patient_name,
            blood_group,
            hospital,
            phone,
            priority
        ))

        conn.commit()
        conn.close()

        return f"""
        <h2>Emergency Request Submitted Successfully!</h2>

        <p>Patient Name: {patient_name}</p>
        <p>Blood Group: {blood_group}</p>
        <p>Priority: {priority}</p>

        <a href="/">Back to Home</a>
        """

    return render_template("request.html")


# =========================
# VIEW REQUESTS
# =========================

@app.route('/requests')
def view_requests():

    conn = sqlite3.connect("donors.db")
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM requests")

    requests = cursor.fetchall()

    conn.close()

    return render_template(
        "requests.html",
        requests=requests
    )


# =========================
# LOGIN
# =========================

@app.route('/login', methods=['GET', 'POST'])
def login():

    if request.method == "POST":

        username = request.form['username']
        password = request.form['password']

        if username == "admin" and password == "1234":

            session['admin'] = True

            return redirect('/admin')

        else:
            return "Invalid Username or Password"

    return render_template("login.html")


# =========================
# ADMIN DASHBOARD
# =========================

@app.route('/admin')
def admin_dashboard():

    if 'admin' not in session:
        return redirect('/login')

    conn = sqlite3.connect("donors.db")
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM donors")
    donors = cursor.fetchall()

    cursor.execute("SELECT * FROM requests")
    requests = cursor.fetchall()

    conn.close()

    return render_template(
        "admin.html",
        donors=donors,
        requests=requests
    )


# =========================
# LOGOUT
# =========================

@app.route('/logout')
def logout():

    session.pop('admin', None)

    return redirect('/login')


# =========================
# BLOOD COMPATIBILITY
# =========================

@app.route('/compatibility', methods=["GET", "POST"])
def compatibility():

    compatible = []

    if request.method == "POST":

        bg = request.form["blood_group"]

        blood_map = {

            "A+": ["A+", "A-", "O+", "O-"],

            "A-": ["A-", "O-"],

            "B+": ["B+", "B-", "O+", "O-"],

            "B-": ["B-", "O-"],

            "AB+": [
                "A+", "A-",
                "B+", "B-",
                "AB+", "AB-",
                "O+", "O-"
            ],

            "AB-": [
                "A-", "B-",
                "AB-", "O-"
            ],

            "O+": ["O+", "O-"],

            "O-": ["O-"]
        }

        compatible = blood_map.get(bg, [])

    return render_template(
        "compatibility.html",
        compatible=compatible
    )


# =========================
# EXCEL UPLOAD
# =========================

@app.route('/upload', methods=['GET', 'POST'])
def upload():

    if request.method == "POST":

        file = request.files['file']

        if file:

            filename = "donors_upload.xlsx"

            file.save(filename)

            workbook = load_workbook(filename)

            sheet = workbook.active

            conn = sqlite3.connect("donors.db")

            cursor = conn.cursor()

            for row in sheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                name, age, blood_group, phone, location = row

                cursor.execute("""
                    INSERT INTO donors
                    (name, age, blood_group, phone, location)
                    VALUES (?, ?, ?, ?, ?)
                """, (
                    name,
                    age,
                    blood_group,
                    phone,
                    location
                ))

            conn.commit()

            conn.close()

            os.remove(filename)

            return "Excel uploaded successfully! All donors have been added."

    return render_template("upload.html")


# =========================
# RUN APPLICATION
# =========================

if __name__ == "__main__":
    app.run(debug=True)